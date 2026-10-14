"""Tests for `run.py hr` (synthetic workbooks built here; never the real HR files).

  .venv/bin/python -I test_hr.py        (from tools/sap-import)

The end-to-end test needs a Postgres it may create a scratch database on: set SAP_TEST_PG to a URL whose user can
CREATE DATABASE (the docker-compose one: postgres://tierra:tierra@localhost:5432/tierra). Without it, it is skipped.
"""
import argparse
import importlib.util
import os
import tempfile
import unittest
import uuid
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
spec = importlib.util.spec_from_file_location("hr_import", HERE / "hr_import.py")
hr = importlib.util.module_from_spec(spec)
spec.loader.exec_module(hr)


def book(path, sheets):
    """sheets: {title: [rows]}; a row is a list of cell values (None for empty)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    wb.save(path)


def write_folder(folder, *, voyon_net_total="60,259.00", register_office_total=29000, joins_note=False):
    f = Path(folder)
    book(f / "voyon-attendance.xlsx", {"Attendance": [
        ["Employee ID", "Employee Name", "Business Unit", "Department", "Designation", "Reporting Officer", "Check In",
         "Check Out", "Check In Location", "Check Out Location", "Status"],
        ["TF901", "Asha Kumar", "Co", "Admin", "Director", None, "", "", None, None, "Not Yet Reported"],
        ["TF902", "Bala Nair", "Co", "Production", "Production Assistant", "Asha Kumar", "08:01 AM", "", "Device1", None, "Present"],
        ["TF903", "Chitra Das", "Co", "QC", "Executive QA", "Asha Kumar", "", "", None, None, "Leave"],
        ["TF905", "Revathy K", "Co", "Production", "Production Assistant", None, "", "", None, None, "Present"],
    ]})
    pay_head = ["Employee ID", "Employee", "Business Unit", "Department", "Designation", "Category", "Pay Days", "Unpaid Days",
                "Basic Pay(A)", "Gross(A)", "Basic Pay(E)", "Gross(E)", "Peeling Incentive(VA)", "Total Earning",
                "LOP(VD)", "TDS(VD)", "PF(SD)", "ESI(SD)", "Total Deductions", "Total Deduction Without LOP", "Net Salary", "Remarks"]
    book(f / "payroll-voyon-2026-02.xlsx", {"Attendance": [
        pay_head,
        ["TF901", "Asha Kumar", "Co", "Admin", "Director", "General", 28, 0, 30000, 30000, 30000, 30000, "", 30000, "", 1000, "", "", 1000, 1000, 29000, None],
        ["TF902", "Bala Nair", "Co", "Production", "Production Assistant", "General", 28, 0, 15000, 18000, 15000, 18000, 4000, 22000, "", "", 1800, "", 1800, 1800, 20200, None],
        ["TF903", "Chitra Das", "Co", "QC", "Executive QA", "General", 26, 2, 12000, 12000, 11143, 11143, "", 11143, 857, "", "", 84, 941, 84, 11059, None],
        [None] * 8 + ["57,000.00", "60,000.00", "56,143.00", "59,143.00", "4,000.00", "63,143.00", "857.00", "1,000.00",
                      "1,800.00", "84.00", "3,741.00", "2,884.00", voyon_net_total, None],
    ]})
    office = [["TIERRA TEST PVT LTD"], ["SALARY REGISTER  February  2026"], [None, 1, 2, 3],
              ["SL No.", "Emp Code", "EMPLOYEE NAME", "DESIGNATION", "Location", "DOJ", "No.of days paid", "BASIC", None, None,
               None, None, None, None, None, None, None, "Total Remuneration", "Deductions", None, "TOTAL DEDUCTIONS", "NET SALARY",
               "Net Salary For January", "ESI Applicability"],
              [None, None, None, None, None, None, None, "BASIC", "DA", "HRA", "Conveyance", "spl all", "Gross Salary",
               "Gross Actual", "Off day work", "Overtime", "Arrears", None, "PF 12%", "TDS", None, None, None, None],
              [1, "T901", "Asha Kumar", "Director", "Plant", datetime(2023, 4, 1), 28, 30000, 0, 0, 0, None, 30000, 30000, 0, 0,
               None, 30000, 0, 1000, 1000, 29000, 29000, "NO"],
              [None] * 21 + [register_office_total]]
    peel = [["TIERRA TEST PVT LTD"], ["SALARY REGISTER February 26- Peeling"], [],
            ["SL No.", "Emp Code", "EMPLOYEE NAME", "DESIGNATION", "Location", "DOJ", "No.of days paid", "BASIC", None, None, None,
             None, None, None, None, "Attendance Incentive ", None, None, None, "Total Remuneration", "Deductions", None,
             "TOTAL DEDUCTIONS", "NET SALARY", "Net Salary"],
            [None, None, None, None, None, None, None, "BASIC", "DA", "HRA", "Conveyance", "Spl.Allowance", "Gross Salary",
             "Gross Actual", "Off day work", None, "Peeling Incentive", "Production Incentive ", None, None, "PF 12% ", "ESI",
             None, None, None],
            [1, "TFL902", "Bala Nair", "Production Assistant", "Plant", datetime(2012, 3, 29), 28, 10000, 4000, 1000, 0, 0, 15000,
             14100, 0, 1000, 3000, 0, None, 18100, 1800, 0, 1800, 16300, 16300],
            [2, "T906", "Revathi K", "Production Assistant", "Plant", None, 20, 10000, 0, 0, 0, 0, 10000, 6667, 0, 0, 0, 0, None,
             6667, 0, 50, 50, 6617, 6617],
            [3, "#REF!", "#REF!", None, None, None, 0, 10000, 0, 0, 0, 0, 10000, 0, 0, 0, 0, 0, None, 0, 0, 0, 0, 0, 0],
            [None] * 23 + [22917]]
    promo = [["TIERRA TEST PVT LTD"], ["Sales Promotion  REGISTER  February  26"], [],
             ["SL No.", "Emp Code", "EMPLOYEE NAME", "DOJ", "No.of days paid", "BASIC", None, None, None, None, None, None,
              "Total Remuneration", None, "NET SALARY"],
             [None, None, None, None, None, "BASIC", "Gross Salary", "Gross Actual", "Off Day Work", "Overtime ",
              "Performance incentive ", "Production Incentive", None, "Salary After Statutory Deductions", None],
             [1, "T904", "Dev Menon", "01.05.25", 25, 12000, 12000, 10000, 0, 0, None, 0, 10000, 10000, 10000],
             [None] * 14 + [10000]]
    attn_head = [" ", "S.No.", "Emp. Code", "Name of Employee", "Location"] + list(range(1, 29)) + ["29", "30", "31"] + \
        ["No. of Working Days", "Sl", "EL", "CL", "CO+", "CO-", "Other Leaves", "PH", "W/Off", "LWP", "Absent ", "OT Hours",
         "shift1", "shift2", "shift3", "Tot. Days in a Month", "Dyas Paid", "EL Bal", "CL Bal", "Offdays"]

    def attn_row(code, name, marks, working, cl, wof, lwp, offdays):
        days = [marks.get(d) for d in range(1, 32)]
        return [None, "1", code, name, "Plant"] + days + [working, 0, 0, cl, 0, None, None, 0, wof, lwp, None, None, None, None,
                                                          None, 28, None, None, None, offdays]
    attn = [[], [None, "Test Co", None, None, None, "Use only following abbreviations"],
            [None, "Attendance Sheet for the Month : ", None, None, datetime(2026, 2, 1), "P", "Present"],
            [None, "Location", None, "  :  ", "Plant", "OD", "On Duty"], [None, "STAFF"], attn_head,
            [None, None, None, None, None, None, None, None, None, None, None, None, "S"],
            attn_row("T901", "ASHA KUMAR", {2: "P", 3: "P", 4: "W/OFF"}, 2, 0, 1, 0, 0),
            attn_row("TFL902", "BALA NAIR", {2: "P", 3: "LWP", 4: "CL", 5: "Offday"}, 2, 1, 0, 1, 1),
            attn_row("T906", "REVATHI K", {2: "p", 3: "P"}, 2, 0, 0, 0, 0)]
    book(f / "salary-register-2026-02.xlsx", {"ATTN-ADOOR": attn, "Salary-Office&Admin": office, "Salary-Peeling": peel,
                                              "Sales Promotion ": promo})
    book(f / "temp-wages-2026-02.xlsx", {"Sheet1": [
        ["Temporary Wages 01.0.25 to 31.01.25"],
        ["Sl", "Name", "Days", "Wages ", "PeelingIncentive", "Attendance inentive", "Total", "ESI", "NET", 0.0325, None,
         "SLAES PROMOTION", None, "Days", "wages", "Peeling Incentive", "ATTENDANCE BASE INCENTIVE", "TOTAL", "ADV", "NET "],
        [None] * 11 + [1, "Gita Promo", 20, 8000, 0, None, 8000, 1000, 7000],
        [1, "Ela Temp", 26, 10400, 600, 1300, 12300, 93, 12207, 400],
        [None, None, None, 10400, 600, 1300, 12300, 93, 12207, 400] + [None] * 4 + [8000, 0, None, 8000, 1000, 7000],
    ]})
    book(f / "leave-requests.xlsx", {"Leave Request": [
        ["Leave Request"], [],
        ["Transaction Date", "Employee ID", "Employee", "Leave Type", "From Date", "To Date", "Number of Leave Day(s)",
         "Reason For Leave", "Requested By", "Status"],
        [datetime(2026, 2, 2, 4), "TF902", "Bala Nair", "Casual Leave", datetime(2026, 2, 4), datetime(2026, 2, 4), "1", None, "TF901", "Approved"],
        [datetime(2026, 2, 9, 5), "TF903", "Chitra Das", "Unpaid Leave", datetime(2026, 2, 10), datetime(2026, 2, 11), "2", "fever", "TF901", "Approved"],
    ]})
    dates = [datetime(2026, 2, 2), datetime(2026, 2, 3)]
    book(f / "peeling-register-2026-02.xlsx", {
        "Peeling Banana": [["Banana Peeling Statement"], [], ["Employee Name", "Date of Joining"],
                           ["1st Shift", None] + dates, ["Bala Nair", None, 210, 170], ["Ela Temp", datetime(2024, 1, 1), 160, 100],
                           ["Operators ", None], ["Total Infeed", None, 370, 270]],
        "Incentive": [["Peeling Incentive"], ["Employee Name"] + dates, ["Bala Nair", 157.5, 42.5], ["Ela Temp", 120, 0],
                      ["Operators ", 0, 0], ["Total", 277.5, 42.5]],
        "Incentive summary": [["Temporaray ", "Total", None, "Permanent", "Total"], [], ["Ela Temp", 120, None, "Bala Nair", 200],
                              ["Total", 120, None, None, 200]],
    })
    head = ["pay_date", "month", "name", "name_key", "amount", "purpose", "cycle", "sap_docnum", "sap_comment", "payment_total",
            "source", "source_path"]
    pdf = "sap_attachments/Attachments/Ewaybill/Payment-Aproval 03 MAR.pdf"
    book(f / "payroll-series-2025-12-to-2026-03.xlsx", {
        "Payroll lines": [
            head,
            [datetime(2026, 3, 3), "2026-03", "Asha Kumar", "ashakumar", 20000, "Jan Salary", "Monthly salary", 501, "Paid salary", 36300, "PDF text", pdf],
            [datetime(2026, 3, 3), "2026-03", "Bala (Asha Kumar)", "bala", 16000, "Feb 26 Salary adv", "Salary advance", 501, "Paid salary", 36300, "PDF text", pdf],
            [datetime(2026, 3, 3), "2026-03", "(not itemised – see sheet)", None, 300, "Paid salary", "Monthly salary", 501, "Paid salary", 36300, "PDF text – balancing line", pdf],
            [datetime(2026, 3, 5), "2026-03", "Salary Advance paid to Ela Temp", "ela", 500, "Salary advance paid/Ela Temp-500,Chitra Das-500", "Salary advance", 502,
             "Salary advance paid/Ela Temp-500,Chitra Das-500", 1000, "pairs typed into SAP comment", "sap_backup/parquet/OVPM.parquet DocNum 502"],
            [datetime(2026, 3, 5), "2026-03", "Chitra Das", "chitra", 500, "Salary advance paid/Ela Temp-500,Chitra Das-500", "Salary advance", 502,
             "Salary advance paid/Ela Temp-500,Chitra Das-500", 1000, "pairs typed into SAP comment", "sap_backup/parquet/OVPM.parquet DocNum 502"],
            [datetime(2026, 3, 6), "2026-03", "(not itemised – see sheet)", None, 990, "Lunch for TN Peelers", "TN peeling (weekly, piece-rate)", 503,
             "Lunch for TN Peelers", 990, "SAP payment only (no sheet)", "sap_backup/parquet/OVPM.parquet DocNum 503"],
            [datetime(2026, 3, 9), "2026-03", "(not itemised – see sheet)", None, 1200, "TDS on salary", "TDS on salary", 504, "TDS on salary", 1200,
             "SAP payment only (no sheet)", "sap_backup/parquet/OVPM.parquet DocNum 504"],
        ],
        "Monthly totals": [["Month", "Total"], ["2026-03", "=SUM(1)"]],
    })
    book(f / "labour-cost.xlsx", {"Aug": [["rates"], [1, 2, 3]]})
    if joins_note:
        (f / "notes.txt").write_text("not a workbook")


def args_for(folder, joins, pg=None, gates="/nonexistent"):
    return argparse.Namespace(dir=str(folder), joins=str(joins), month=None, dry_run=pg is None, pg=pg, gates=gates)


class Helpers(unittest.TestCase):
    def test_cells_and_names(self):
        self.assertEqual(hr.num("1,23,456.50"), Decimal("123456.50"))
        self.assertIsNone(hr.num("#REF!"))
        self.assertIsNone(hr.num(False))
        self.assertEqual(hr.name_key("Orvetha . T"), "orvetha t")
        self.assertEqual(hr.loose_key("Zarina R"), hr.loose_key("Zharina R"))
        self.assertEqual(hr.compact_key("Velora Kumari P"), hr.compact_key("Velorakumari P"))
        self.assertEqual(hr.compact_key("Quenly G"), hr.compact_key("Quenli G"))
        self.assertEqual(hr.month_in("SALARY REGISTER June  26- Peeling").isoformat(), "2026-06-01")
        self.assertEqual(hr.month_in("hr-payroll-voyon-2026-06").isoformat(), "2026-06-01")
        self.assertIsNone(hr.month_in("Attendance Sheet"))

    def test_payment_names_and_months(self):
        self.assertEqual(hr.payee_name("Salary paid to sreelaja"), ("sreelaja", None))
        self.assertEqual(hr.payee_name("Sharook (Noufal)"), ("Sharook", "Noufal"))
        self.assertEqual(hr.payee_name("Sep & Oct Abiya"), ("Abiya", None))
        self.assertEqual(hr.payee_name("NO - Midhun"), ("Midhun", None))
        self.assertEqual(hr.payee_name("(not itemised – see sheet)"), (None, None))
        paid = datetime(2026, 3, 3).date()
        self.assertEqual(hr.salary_month("July Salary", paid).isoformat(), "2025-07-01")
        self.assertEqual(hr.salary_month("Mar salary", paid).isoformat(), "2026-03-01")
        self.assertEqual(hr.salary_month("Apr salary", paid).isoformat(), "2025-04-01")      # not ahead of the payment
        self.assertEqual(hr.salary_month("Apr salary adv", paid).isoformat(), "2026-04-01")  # an advance may be
        self.assertEqual(hr.salary_month("Dec 24 salary", paid).isoformat(), "2024-12-01")
        self.assertIsNone(hr.salary_month("Sep 24 and Oct 24 salary", paid))
        self.assertIsNone(hr.salary_month("net pay", paid))

    def test_peeling_slabs(self):
        self.assertEqual(hr.slab_incentive(Decimal(170), "permanent"), Decimal("42.50"))
        self.assertEqual(hr.slab_incentive(Decimal(210), "permanent"), Decimal("157.50"))
        self.assertEqual(hr.slab_incentive(Decimal("200.5"), "permanent"), Decimal(0))  # the register's own gap
        self.assertEqual(hr.slab_incentive(Decimal(150), "temporary"), Decimal("112.50"))
        self.assertEqual(hr.slab_incentive(Decimal(149), "temporary"), Decimal(0))


class Build(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.folder = Path(self.tmp.name) / "hr"
        self.folder.mkdir()
        self.joins = Path(self.tmp.name) / "joins.yaml"
        self.joins.write_text("joins:\n  - {register: T901, voyon: TF901}\n")

    def tearDown(self):
        self.tmp.cleanup()

    def gates(self, results):
        return {name: (ok, detail) for name, _, ok, detail in results}

    def test_reads_every_file_and_gates_pass(self):
        write_folder(self.folder, joins_note=True)
        ds, report, results = hr.build(args_for(self.folder, self.joins))
        text = "\n".join(report)
        self.assertIn("skipped           labour-cost.xlsx", text)
        failed = [(n, d) for n, kind, ok, d in results if ok is False and kind == "hard"]
        self.assertEqual(failed, [])
        g = self.gates(results)
        self.assertEqual(g["Voyon 2026-02: every column adds up to the file's totals row"][0], True)
        self.assertIn("1 row(s) with a broken code (#REF!)", g["register 2026-02 Salary-Peeling: net = totals row"][1])
        self.assertIn("P 5", g["attendance 2026-02: grid = the sheet's own day counts"][1])
        self.assertEqual(g["peeling 2026-02 permanent incentive = the summary sheet"][0], True)

        emp = {e["employee_id"]: e for e in ds["employees"]}
        self.assertEqual(emp["TF901"]["code_register"], "T901")
        self.assertEqual(emp["TF901"]["matched_by"], "joins_file")
        self.assertEqual(emp["TF902"]["code_register"], "TFL902")      # same name
        self.assertEqual(emp["TF905"]["code_register"], "T906")        # Revathy / Revathi
        self.assertEqual(emp["TF905"]["matched_by"], "loose_name")
        self.assertIn("T904", emp)                                    # register only
        self.assertEqual(emp["TF902"]["reporting_officer_id"], "TF901")
        self.assertEqual(emp["TF902"]["punch_status"], "Present")

        runs = {r["run_id"]: r for r in ds["runs"]}
        self.assertEqual(runs["voyon-2026-02"]["net"], Decimal("60259.00"))  # 29000 + 20200 + 11059
        self.assertEqual(runs["register-2026-02"]["headcount"], 4)
        self.assertEqual(runs["temp_sheet-2026-02"]["net"], Decimal("19207.00"))
        cats = sorted((l["category"], l["net"]) for l in ds["lines"] if l["run_id"] == "register-2026-02")
        self.assertEqual(cats, [("office_admin", Decimal("29000.00")), ("peeling", Decimal("6617.00")),
                                ("peeling", Decimal("16300.00")), ("sales_promotion", Decimal("10000.00"))])
        bala = next(l for l in ds["lines"] if l["run_id"] == "register-2026-02" and l["code"] == "TFL902")
        self.assertEqual((bala["attendance_incentive"], bala["peeling_incentive"], bala["total_earning"]),
                         (Decimal("1000.00"), Decimal("3000.00"), Decimal("18100.00")))
        temps = {(t["block"], t["name"]) for t in ds["temp_workers"]}
        self.assertEqual(temps, {("temporary", "Ela Temp"), ("sales_promotion_temp", "Gita Promo")})
        codes = sorted(a["code"] for a in ds["attendance"])
        self.assertEqual(codes, ["CL", "LWP", "Offday", "P", "P", "P", "P", "P", "W/OFF"])
        self.assertEqual(sum(p["incentive"] for p in ds["peeling"]), Decimal("320.00"))
        self.assertEqual({p["worker_name"]: p["employee_id"] for p in ds["peeling"]}, {"Bala Nair": "TF902", "Ela Temp": "TMP-001"})
        self.assertEqual(len(ds["leaves"]), 2)

        self.assertIn("salary payments 2026-03-03 .. 2026-03-09: 3 payments, 38,500.00", text)
        self.assertIn("left out (not wages): 1 payments, 990.00", text)
        self.assertTrue(g["salary payments (payroll-series-2025-12-to-2026-03.xlsx): lines = the SAP payment total"][0])
        pays = {p["sap_doc_num"]: p for p in ds["payments"]}
        self.assertEqual(sorted(pays), [501, 502, 504])
        self.assertEqual((pays[501]["cycle"], pays[501]["itemised"], pays[501]["list_file"]),
                         ("monthly_salary", False, "Payment-Aproval 03 MAR.pdf"))
        self.assertEqual((pays[502]["cycle"], pays[502]["itemised"], pays[502]["list_file"]), ("advance", True, None))
        self.assertEqual([l["cycle"] for l in ds["payment_lines"] if l["sap_doc_num"] == 501], ["monthly_salary", "advance", "monthly_salary"])
        got = [(l["sap_doc_num"], l["name"], l["name_note"], l["employee_id"], l["amount"], str(l["for_month"]), l["how_read"])
               for l in ds["payment_lines"]]
        self.assertEqual(got, [
            (501, "Asha Kumar", None, "TF901", Decimal("20000.00"), "2026-01-01", "pdf"),
            (501, "Bala", "Asha Kumar", None, Decimal("16000.00"), "2026-02-01", "pdf"),   # an advance on a salary list   # a first name alone does not link
            (501, None, None, None, Decimal("300.00"), "None", "balancing"),
            (502, "Ela Temp", None, "TMP-001", Decimal("500.00"), "None", "sap_comment"),
            (502, "Chitra Das", None, "TF903", Decimal("500.00"), "None", "sap_comment"),
            (504, None, None, None, Decimal("1200.00"), "None", "none"),
        ])

    def test_payment_lines_that_miss_the_sap_total_fail_their_gate(self):
        write_folder(self.folder)
        wb = openpyxl.load_workbook(self.folder / "payroll-series-2025-12-to-2026-03.xlsx")
        wb["Payroll lines"]["E2"] = 19000
        wb.save(self.folder / "payroll-series-2025-12-to-2026-03.xlsx")
        _, _, results = hr.build(args_for(self.folder, self.joins))
        ok, detail = self.gates(results)["salary payments (payroll-series-2025-12-to-2026-03.xlsx): lines = the SAP payment total"]
        self.assertFalse(ok)
        self.assertIn("501 of 2026-03-03: 35300.00 vs 36300.00", detail)

    def test_a_total_that_does_not_add_up_fails_its_gate(self):
        write_folder(self.folder, voyon_net_total="61,000.00", register_office_total=28000)
        _, _, results = hr.build(args_for(self.folder, self.joins))
        g = self.gates(results)
        self.assertFalse(g["Voyon 2026-02: every column adds up to the file's totals row"][0])
        self.assertIn("Net Salary 60259 vs 61000.00", g["Voyon 2026-02: every column adds up to the file's totals row"][1])
        self.assertFalse(g["register 2026-02 Salary-Office&Admin: net = totals row"][0])


@unittest.skipUnless(os.environ.get("SAP_TEST_PG"), "set SAP_TEST_PG to run the Postgres test")
class EndToEnd(unittest.TestCase):
    def setUp(self):
        import psycopg
        self.psycopg = psycopg
        self.name = f"hr_test_{uuid.uuid4().hex[:8]}"
        base = os.environ["SAP_TEST_PG"]
        with psycopg.connect(base, autocommit=True) as pg:
            pg.execute(f"CREATE DATABASE {self.name}")
        self.url = base.rsplit("/", 1)[0] + "/" + self.name
        self.tmp = tempfile.TemporaryDirectory()
        self.folder = Path(self.tmp.name) / "hr"
        self.folder.mkdir()
        self.joins = Path(self.tmp.name) / "joins.yaml"
        self.joins.write_text("joins:\n  - {register: T901, voyon: TF901}\n")
        self.gates_file = Path(self.tmp.name) / "gates.yaml"
        self.gates_file.write_text(
            "hr_checks:\n  - name: voyon net\n    sql: SELECT sum(net) FROM {s}.payroll_lines WHERE run_id = 'voyon-2026-02'\n"
            "    expect: 60259\n")

    def tearDown(self):
        with self.psycopg.connect(os.environ["SAP_TEST_PG"], autocommit=True) as pg:
            pg.execute(f"DROP DATABASE IF EXISTS {self.name} WITH (FORCE)")
        self.tmp.cleanup()

    def test_loads_swaps_and_keeps_hr_on_a_failed_gate(self):
        write_folder(self.folder, voyon_net_total="60,259.00")
        self.assertEqual(hr.run(args_for(self.folder, self.joins, self.url, self.gates_file)), 0)
        with self.psycopg.connect(self.url) as pg:
            self.assertEqual(pg.execute("SELECT count(*), sum(net) FROM hr.payroll_lines WHERE run_id = 'voyon-2026-02'").fetchone(),
                             (3, Decimal("60259.00")))
            self.assertEqual(pg.execute("SELECT count(*) FROM hr.attendance_days").fetchone()[0], 9)
            self.assertIsNone(pg.execute("SELECT to_regnamespace('hr_next')").fetchone()[0])
            self.assertEqual(pg.execute("SELECT count(*), sum(amount) FROM hr.salary_payment_lines").fetchone(), (6, Decimal("38500.00")))
        # a run whose totals do not add up leaves the loaded hr alone and keeps hr_next to look at
        write_folder(self.folder, voyon_net_total="1.00")
        self.assertEqual(hr.run(args_for(self.folder, self.joins, self.url, self.gates_file)), 1)
        with self.psycopg.connect(self.url) as pg:
            self.assertEqual(pg.execute("SELECT count(*) FROM hr.payroll_lines").fetchone()[0], 9)
            self.assertIsNotNone(pg.execute("SELECT to_regnamespace('hr_next')").fetchone()[0])


if __name__ == "__main__":
    unittest.main(verbosity=2)
